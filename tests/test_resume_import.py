"""Synthetic resume documents, private owner review, and atomic source provenance."""

from __future__ import annotations

import io
import json
import sqlite3
import subprocess
import warnings
import zipfile
from datetime import timedelta
from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from internship_pipeline.app import create_app
from internship_pipeline.assessments import current_identity
from internship_pipeline.models import FetchResult, SourceJob, utcnow
from internship_pipeline.profile_settings import (
    Fact,
    Preferences,
    Profile,
    RevisionConflict,
    SaveSettings,
)
from internship_pipeline.profiles import imports
from internship_pipeline.profiles.extract import MAX_PDF_CONTENT, MAX_UPLOAD, DocumentError, extract
from internship_pipeline.profiles.imports import ResumeImports, Review, Selection, extract_bounded
from internship_pipeline.storage import Store

SOURCE = [
    "Synthetic Candidate",
    "synthetic@example.test",
    "Education",
    "Example University",
    "Bachelor of Science in Computing",
    "Skills",
    "Python, SQL",
    "Experience",
    "Built an example reporting tool.",
    "Projects",
    "Created a synthetic data dashboard.",
    "Interests",
    "Rust",
]


def docx_bytes(lines: list[str] = SOURCE) -> bytes:
    document = Document()
    for text in lines:
        document.add_paragraph(text)
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


def pdf_bytes(*, encrypted: bool = False, pages: int = 1, text: bool = True) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            content = (
                "BT /F1 12 Tf 50 740 Td 16 TL "
                + " ".join(f"({line}) Tj T*" for line in SOURCE)
                + " ET"
            )
            stream.set_data(content.encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-password")
    data = io.BytesIO()
    writer.write(data)
    return data.getvalue()


@pytest.fixture
def service(tmp_path: Path) -> ResumeImports:
    return ResumeImports(Store(tmp_path / "state.sqlite3"))


def reviewed(draft: dict, **changes) -> Review:
    selections = [
        Selection.model_validate({key: value for key, value in item.items() if key != "selected"})
        for item in draft["suggestions"]
        if item["selected"]
    ]
    return Review(
        expected_revision=draft["expected_revision"],
        selections=selections,
        confirmed=True,
        **changes,
    )


@pytest.mark.parametrize(
    "kind,data", [("pdf", pdf_bytes()), ("docx", docx_bytes())], ids=["pdf", "docx"]
)
def test_real_extraction_proposes_literal_useful_draft(kind: str, data: bytes) -> None:
    lines = extract_bounded(data, kind)
    assert [line["text"] for line in lines] == SOURCE
    assert lines[0]["location"].startswith("Page" if kind == "pdf" else "Paragraph")
    proposed = imports.suggestions(lines)
    assert [item["kind"] for item in proposed if item["selected"]] == [
        "name",
        "email",
        "education",
        "skill",
        "experience",
        "project",
    ]
    assert next(item for item in proposed if item["kind"] == "skill")["skills"] == ["Python", "SQL"]
    assert not proposed[-1]["selected"]  # Rust is an interest, not an acquired skill.


def test_docx_tables_headers_and_source_order() -> None:
    document = Document()
    document.add_paragraph("Synthetic Candidate")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "SQL"
    document.sections[0].header.paragraphs[0].text = "Synthetic header"
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
    data = io.BytesIO()
    document.save(data)
    lines = extract(data.getvalue(), "docx")
    assert [item["text"] for item in lines] == [
        "Synthetic Candidate",
        "Python | SQL",
        "Synthetic header",
        "Synthetic footer",
    ]
    assert lines[1]["location"].startswith("Table")


def test_decoded_pdf_page_content_bound() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(b" " * (MAX_PDF_CONTENT + 1))
    page[NameObject("/Contents")] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    with pytest.raises(DocumentError, match="page content exceeds 2 MiB"):
        extract_bounded(data.getvalue(), "pdf")


@pytest.mark.parametrize(
    "data,kind,message",
    [
        (b"", "pdf", "nonempty"),
        (b"x" * (MAX_UPLOAD + 1), "docx", "5 MiB"),
        (b"not pdf", "pdf", "not a PDF"),
        (b"not docx", "docx", "not a DOCX"),
        (b"anything", "tex", "Only PDF"),
        (pdf_bytes(encrypted=True), "pdf", "Encrypted"),
        (pdf_bytes(pages=31), "pdf", "30 pages"),
        (pdf_bytes(text=False), "pdf", "Image-only"),
        (docx_bytes([]), "docx", "Image-only"),
        (docx_bytes(["x" * 2001]), "docx", "2,000"),
        (docx_bytes(["line"] * 501), "docx", "500 text lines"),
        (docx_bytes(["x" * 250] * 450), "docx", "100,000"),
        (bytes.fromhex("d0cf11e0"), "docx", "Encrypted or older"),
    ],
)
def test_document_rejections(data: bytes, kind: str, message: str) -> None:
    with pytest.raises(DocumentError, match=message):
        extract(data, kind)


@pytest.mark.parametrize(
    "name,data,compression,message",
    [
        ("word/document.xml", b"x" * 10000, zipfile.ZIP_DEFLATED, "compression"),
        ("word/document.xml", b"<!DOCTYPE fake>", zipfile.ZIP_STORED, "XML entities"),
        ("word/vbaProject.bin", b"macro", zipfile.ZIP_STORED, "Macros"),
        ("word/embeddings/object.bin", b"object", zipfile.ZIP_STORED, "embedded objects"),
    ],
)
def test_docx_archive_hazards(name: str, data: bytes, compression: int, message: str) -> None:
    value = io.BytesIO()
    with zipfile.ZipFile(value, "w", compression=compression) as archive:
        archive.writestr(name, data)
    with pytest.raises(DocumentError, match=message):
        extract(value.getvalue(), "docx")


@pytest.mark.parametrize(
    "hazard,message",
    [
        ("parts", "512-part"),
        ("expanded", "20 MiB"),
        ("duplicates", "duplicate"),
        ("encrypted", "Encrypted"),
    ],
)
def test_archive_metadata_bounds(hazard: str, message: str) -> None:
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        if hazard == "parts":
            for index in range(513):
                archive.writestr(f"part-{index}", b"x")
        elif hazard == "expanded":
            archive.writestr("word/document.xml", b"x" * (20 * 1024 * 1024 + 1))
        elif hazard == "duplicates":
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                archive.writestr("word/document.xml", b"one")
                archive.writestr("word/document.xml", b"two")
        else:
            archive.writestr("word/document.xml", b"private")
    payload = bytearray(data.getvalue())
    if hazard == "encrypted":
        # Mark both the local header and central directory as encrypted.
        local = payload.index(b"PK\x03\x04") + 6
        central = payload.index(b"PK\x01\x02") + 8
        payload[local] |= 1
        payload[central] |= 1
    with pytest.raises(DocumentError, match=message):
        extract(bytes(payload), "docx")


def test_draft_capacity_and_literal_labeled_skills(service: ResumeImports) -> None:
    facts = [
        Fact(id=f"manual-{index}", text=f"Supported manual fact {index}.") for index in range(99)
    ]
    service.settings.save(
        SaveSettings(expected_revision=0, profile=Profile(facts=facts), preferences=Preferences())
    )
    draft = service.upload(
        docx_bytes(
            [
                "Synthetic Candidate",
                "Skills",
                "Languages: Python, SQL",
                "Experience",
                "Another supported fact.",
            ]
        ),
        "docx",
        1,
    )
    selected = [item for item in draft["suggestions"] if item["selected"]]
    assert len(selected) == 2
    assert selected[-1]["skills"] == ["Python", "SQL"]
    assert not draft["suggestions"][-1]["selected"]


def test_bounded_subprocess_timeout_failure_and_busy(monkeypatch: pytest.MonkeyPatch) -> None:
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("extract", 12)

    monkeypatch.setattr(imports.subprocess, "run", timeout)
    with pytest.raises(DocumentError, match="too long"):
        extract_bounded(b"x", "pdf")
    monkeypatch.setattr(
        imports.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess("extract", -9, b"")
    )
    with pytest.raises(DocumentError, match="resource limits"):
        extract_bounded(b"x", "pdf")
    assert imports._PROCESS_SLOT.acquire(blocking=False)
    try:
        with pytest.raises(HTTPException) as error:
            extract_bounded(b"x", "pdf")
        assert error.value.status_code == 429
    finally:
        imports._PROCESS_SLOT.release()


def test_review_preview_confirm_stable_ids_provenance_and_restart(service: ResumeImports) -> None:
    draft = service.upload(docx_bytes(), "docx", 0)
    assert service.settings.read().revision == 0
    assert service.current() == {"import": None}
    body = reviewed(draft)
    prepared, evidence = service.prepare(draft["id"], body)
    saved = service.confirm(draft["id"], body)
    assert saved.profile == prepared.profile
    assert saved.profile.name == SOURCE[0]
    assert saved.profile.email == SOURCE[1]
    assert saved.profile.education[0].degree == SOURCE[4]
    assert saved.profile.requires_sponsorship is None
    assert saved.profile.education[0].graduation_date is None
    assert saved.profile.facts[0].skills == ["Python", "SQL"]
    assert "Rust" not in {skill for fact in saved.candidate().facts for skill in fact.skills}
    assert ResumeImports(Store(service.store.path)).current()["import"]["evidence"] == evidence
    with service.store.connection() as connection:
        assert connection.execute("SELECT approved_revision FROM resume_imports").fetchone()[0] == 1
        assert (
            json.loads(connection.execute("SELECT lines FROM resume_imports").fetchone()[0])[0][
                "text"
            ]
            == SOURCE[0]
        )
    with pytest.raises(RevisionConflict, match="already confirmed"):
        service.confirm(draft["id"], body)


def test_replace_preserves_manual_edited_facts_and_preferences(service: ResumeImports) -> None:
    first = service.upload(docx_bytes(), "docx", 0)
    saved = service.confirm(first["id"], reviewed(first))
    edited = saved.profile.model_copy(deep=True)
    edited.facts[1].text = "Manually refined supported reporting tool claim."
    manual = Fact(
        id="manual", kind="project", status="confirmed", text="A manually confirmed project."
    )
    edited.facts.append(manual)
    prefs = Preferences.model_validate({"soft": {"skills": ["Rust"]}})
    saved = service.settings.save(
        SaveSettings(expected_revision=1, profile=edited, preferences=prefs)
    )
    imported_id = edited.facts[0].id
    replacement = service.upload(
        docx_bytes(["Synthetic Candidate", "Skills", "Python, Go"]), "docx", 2
    )
    removable = {item["id"] for item in replacement["removable"]}
    assert imported_id in removable
    assert manual.id not in removable and edited.facts[1].id not in removable
    body = reviewed(replacement, remove_ids=[imported_id])
    after, _ = service.prepare(replacement["id"], body)
    assert manual in after.profile.facts and edited.facts[1] in after.profile.facts
    assert after.preferences == prefs
    confirmed = service.confirm(replacement["id"], body)
    assert confirmed.profile == after.profile
    assert imported_id not in {item.id for item in confirmed.profile.facts}
    assert manual.id in {item.id for item in confirmed.profile.facts}


def test_reimport_original_claim_preserves_edited_entry_id(service: ResumeImports) -> None:
    draft = service.upload(docx_bytes(), "docx", 0)
    saved = service.confirm(draft["id"], reviewed(draft))
    original_id = saved.profile.facts[1].id
    original_text = saved.profile.facts[1].text
    saved.profile.facts[1].text = "Manually edited supported claim."
    service.settings.save(
        SaveSettings(expected_revision=1, profile=saved.profile, preferences=saved.preferences)
    )
    replacement = service.upload(docx_bytes(), "docx", 2)
    prepared, _ = service.prepare(replacement["id"], reviewed(replacement))
    confirmed = service.confirm(replacement["id"], reviewed(replacement))
    assert prepared.profile == confirmed.profile
    assert (
        next(f for f in confirmed.profile.facts if f.id == original_id).text
        == "Manually edited supported claim."
    )
    restored = next(f for f in confirmed.profile.facts if f.text == original_text)
    assert restored.id != original_id


def test_stale_review_and_expiry_do_not_save(service: ResumeImports) -> None:
    draft = service.upload(docx_bytes(), "docx", 0)
    service.settings.save(
        SaveSettings(expected_revision=0, profile=Profile(), preferences=Preferences())
    )
    with pytest.raises(RevisionConflict, match="Settings changed"):
        service.confirm(draft["id"], reviewed(draft))
    second = service.upload(docx_bytes(), "docx", 1)
    with service.store.transaction() as connection:
        connection.execute(
            "UPDATE resume_imports SET created_at=? WHERE id=?",
            ((utcnow() - timedelta(days=2)).isoformat(), second["id"]),
        )
    with pytest.raises(HTTPException) as error:
        service.prepare(second["id"], reviewed(second))
    assert error.value.status_code == 404
    assert service.settings.read().revision == 1


def test_import_confirmation_changes_matching_identity(service: ResumeImports) -> None:
    service.store.register_target("synthetic", "company", "{}", "synthetic")
    service.store.ingest("synthetic", FetchResult(), "synthetic")
    posting = SourceJob(
        source="synthetic",
        source_id="one",
        board_id="synthetic",
        company="Example",
        title="Intern",
        description="Synthetic job",
        apply_url="https://example.test/job",
    )
    service.store.ingest("synthetic", FetchResult(jobs=[posting]), "synthetic")
    job = service.store.list_jobs()[0]
    before = service.settings.read().candidate().revision
    with service.store.connection() as connection:
        old_identity = current_identity(connection, job)
    draft = service.upload(docx_bytes(), "docx", 0)
    service.confirm(draft["id"], reviewed(draft))
    with service.store.connection() as connection:
        assert current_identity(connection, job) != old_identity
    assert service.settings.read().candidate().revision != before


def test_source_validation_manual_remove_and_explicit_confirmation(service: ResumeImports) -> None:
    draft = service.upload(docx_bytes(), "docx", 0)
    body = reviewed(draft)
    with pytest.raises(DocumentError, match="Explicitly confirm"):
        service.confirm(draft["id"], body.model_copy(update={"confirmed": False}))
    with pytest.raises(DocumentError, match="protected"):
        service.prepare(draft["id"], body.model_copy(update={"remove_ids": ["manual"]}))
    skill = next(selection for selection in body.selections if selection.kind == "skill")
    skill.skills = ["Invented skill"]
    with pytest.raises(DocumentError, match="exact excerpts"):
        service.prepare(draft["id"], body)
    assert service.settings.read().revision == 0


def test_atomic_provenance_failure_rolls_back_profile(service: ResumeImports) -> None:
    def fail(connection: sqlite3.Connection, revision: int) -> None:
        connection.execute(
            "INSERT INTO resume_imports VALUES('rollback', 'now', 'pdf', 'sha', 0, '[]', ?, '[]')",
            (revision,),
        )
        raise RuntimeError("synthetic persistence failure")

    with pytest.raises(RuntimeError):
        service.settings.save(
            SaveSettings(expected_revision=0, profile=Profile(), preferences=Preferences()),
            record_import=fail,
        )
    assert service.settings.read().revision == 0
    with service.store.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM resume_imports").fetchone()[0] == 0
    service.settings.save(
        SaveSettings(expected_revision=0, profile=Profile(), preferences=Preferences())
    )
    with pytest.raises(RevisionConflict):
        service.settings.save(
            SaveSettings(expected_revision=0, profile=Profile(), preferences=Preferences()),
            record_import=fail,
        )


def test_authenticated_csrf_bounded_upload_preview_save_redaction(tmp_path: Path) -> None:
    app = create_app(tmp_path, origin="http://localhost", static_dir=tmp_path)
    with TestClient(app, base_url="http://localhost") as client:
        upload = "/api/resume-imports/upload?format=docx&expected_revision=0"
        assert client.get("/api/resume-imports/current").status_code == 401
        assert client.post(upload, content=docx_bytes()).status_code == 401
        csrf = client.get("/api/session").json()["csrf"]
        claim = client.post(
            "/api/claim",
            headers={"X-CSRF-Token": csrf},
            json={
                "setup_token": app.state.identity.setup_token(rotate=True),
                "username": "synthetic",
                "password": "synthetic-profile-password",
            },
        )
        headers = {"X-CSRF-Token": claim.json()["csrf"]}
        assert client.post(upload, content=docx_bytes()).status_code == 403
        assert (
            client.post(
                upload, headers={**headers, "Origin": "https://evil.test"}, content=docx_bytes()
            ).status_code
            == 403
        )
        assert (
            client.post(upload, headers=headers, content=b"x" * (MAX_UPLOAD + 1)).status_code == 413
        )
        invalid = client.post(upload, headers=headers, content=b"private-document-marker")
        assert invalid.status_code == 422 and "private-document-marker" not in invalid.text
        response = client.post(upload, headers=headers, content=docx_bytes())
        assert response.status_code == 200, response.text
        draft = response.json()
        body = reviewed(draft).model_dump(mode="json")
        preview = client.post(
            f"/api/resume-imports/{draft['id']}/preview", headers=headers, json=body
        )
        assert preview.status_code == 200
        save = client.post(f"/api/resume-imports/{draft['id']}/confirm", headers=headers, json=body)
        assert save.status_code == 200
        assert save.json()["profile"] == preview.json()["after"]["profile"]
        assert (
            client.post(
                f"/api/resume-imports/{draft['id']}/confirm", headers=headers, json=body
            ).status_code
            == 409
        )
        assert client.get("/api/resume-imports/current").headers["cache-control"] == "no-store"
